from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "Audit_Compass_系统统一名词与属性定义表_V1.0.docx"

BLUE = "1F5FAE"
DARK_BLUE = "1E3553"
MID_BLUE = "4A6F9E"
LIGHT_BLUE = "EAF2FB"
PALE_BLUE = "F6F9FD"
INK = "26384F"
MUTED = "65758A"
BORDER = "DCE6F1"
LIGHT_BORDER = "E8EEF5"
WHITE = "FFFFFF"
GREEN = "2F8576"
RED = "A24F5C"
GOLD = "8A642D"
FONT_CN = "Arial Unicode MS"
FONT_EN = "Arial"
TABLE_WIDTH = 9360
TABLE_INDENT = 120


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=110, bottom=80, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])


def set_table_borders(table, color=BORDER, size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_run_font(run, size=9, bold=False, color=INK, name=FONT_CN):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_EN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_EN)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, 8, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    tail = paragraph.add_run(" 页")
    set_run_font(tail, 8, color=MUTED)


def add_callout(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.line_spacing = 1.15
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), PALE_BLUE)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:color"), BLUE)
    borders.append(left)
    p_pr.append(borders)
    lead = p.add_run(f"{label}  ")
    set_run_font(lead, 9.5, bold=True, color=BLUE)
    body = p.add_run(text)
    set_run_font(body, 9.5, color=INK)


def add_section_heading(doc, number, title, description=None):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.keep_with_next = True
    prefix = p.add_run(f"{number}  ")
    set_run_font(prefix, 15, bold=True, color=BLUE)
    title_run = p.add_run(title)
    set_run_font(title_run, 15, bold=True, color=DARK_BLUE)
    if description:
        sub = doc.add_paragraph()
        sub.paragraph_format.space_before = Pt(0)
        sub.paragraph_format.space_after = Pt(6)
        sub.paragraph_format.keep_with_next = True
        run = sub.add_run(description)
        set_run_font(run, 9, color=MUTED)


def add_term_table(doc, rows):
    headers = ["标准名称", "名称定义", "使用说明"]
    widths = [2180, 4080, 3100]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, LIGHT_BLUE)
        set_cell_margins(cell, top=95, bottom=95)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header)
        set_run_font(run, 8.3, bold=True, color=DARK_BLUE)
    for row_index, item in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        values = [item[1], item[4], item[5]]
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2:
                set_cell_shading(cell, "FBFCFE")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(value)
            set_run_font(run, 8.5, bold=col_index == 0, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_attribute_table(doc, rows):
    headers = ["属性名称", "属性定义", "使用说明"]
    widths = [2180, 4080, 3100]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, LIGHT_BLUE)
        set_cell_margins(cell, top=95, bottom=95)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header)
        set_run_font(run, 8.3, bold=True, color=DARK_BLUE)
    for row_index, item in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        values = [item[1], item[4], item[5]]
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index % 2:
                set_cell_shading(cell, "FBFCFE")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(value))
            set_run_font(run, 8.4, bold=col_index == 0, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


TERM_SECTIONS = [
    ("01", "平台、项目与工作簿", "定义系统信息架构和用户在平台中的工作边界。", [
        ("SYS-001", "KPMG Audit Compass", "platformName", "平台", "面向审计项目的数据管理、分析、核查与成果导出平台。", "顶层产品名称；仅用于品牌区、登录页和正式报告。", "审计罗盘（非正式简称）"),
        ("SYS-002", "项目", "project", "业务对象", "承载同一审计业务约定、客户信息、成员权限和一个或多个工作簿的顶层容器。", "主要属性：项目编号、项目名称、客户名称、项目状态、项目成员。", "工作簿、客户项目（混指）"),
        ("SYS-003", "工作簿", "workbook", "业务对象", "项目下的一次独立数据分析任务及其数据、参数、结果和导出成果的集合。", "主要属性：工作簿名称、分析期间、报告期间、IPO/PIE属性、重要性水平、备注。", "项目、报告（混指）"),
        ("SYS-004", "工作簿数据分析", "workbookAnalytics", "导航分组", "进入工作簿后，对数据上传、校验、分析、查询和导出功能的统一导航分组。", "仅作为左侧导航分组名称；不作为单独分析结果名称。", "工作簿详情、工作区"),
        ("SYS-005", "数据上传及管理", "dataManagement", "功能模块", "负责原始数据上传、数据质量校验和分析范围管控。", "包含数据上传、数据校验、分析范围管控三个页签。", "数据上传、数据校验（作为模块总称）"),
        ("SYS-006", "资金流水分析", "fundFlowAnalysis", "功能模块", "基于银行流水和相关主档开展账户、经营、特殊交易及明细分析。", "数据源仍称“银行流水”；模块统一称“资金流水分析”。", "银行流水分析"),
        ("SYS-007", "增值税数据分析", "vatAnalysis", "功能模块", "基于销项、进项发票和税务主档开展税号、购销、关联交易及查询分析。", "正式模块名称，不简称“税票分析”。", "增值税分析、税票分析"),
        ("SYS-008", "数据综合分析", "integratedAnalysis", "功能模块", "将银行流水与增值税发票按对手方、月份和账期口径进行客观比较。", "不直接输出异常或风险结论；重点展示总量、账期、构成和明细。", "综合分析"),
        ("SYS-009", "报告及底稿导出", "reportWorkingPaperExport", "功能模块", "统一生成风险评估报告和资金流水核查底稿。", "包含风险评估报告导出、资金流水核查底稿导出两个页签。", "资金流水核查底稿导出（作为模块总称）"),
        ("SYS-010", "分析范围", "analysisScope", "业务概念", "当前工作簿纳入计算的公司、账户、税号、期间和对手方集合。", "由主档、上传数据、筛选条件和范围管控结果共同决定。", "数据范围（未说明范围维度）"),
    ]),
    ("02", "主体与主档", "统一公司、对手方、账户和税务主体的身份表达。", [
        ("ENT-001", "被审计单位公司", "auditEntity", "主体", "纳入当前审计工作簿分析范围的被审计法人或分支机构。", "筛选、表格和导出统一使用；允许多选或“全部”。", "被审计单位、公司"),
        ("ENT-002", "集团内公司", "groupEntity", "主体", "根据被审计单位主档或外部工商关系识别的同一集团范围内公司。", "用于集团内往来识别和含/剔除集团口径计算。", "关联公司（泛称）"),
        ("ENT-003", "对手方", "counterparty", "主体", "与被审计单位发生银行收付款或发票购销关系的个人或组织。", "上位概念；下分客户、供应商、员工、个人、其他等类型。", "交易方、客商（正式字段）"),
        ("ENT-004", "客户", "customer", "对手方类型", "客户主档中有效，或依据销项发票、销售业务关系识别的对手方。", "应保留批准日期、无效日期和销售类别。", "销售方（方向错误）"),
        ("ENT-005", "供应商", "supplier", "对手方类型", "供应商主档中有效，或依据进项发票、采购业务关系识别的对手方。", "应保留批准日期、无效日期和采购类别。", "采购方（方向错误）"),
        ("ENT-006", "员工", "employee", "对手方类型", "员工主档中存在有效任职关系的自然人。", "主要属性：员工编号、职位、入职日期、离职日期、银行账号。", "个人（未区分员工）"),
        ("ENT-007", "个人", "individual", "对手方类型", "以自然人身份出现的对手方，包括员工和其他个人。", "“员工”为其子类；统计时需明确是否含员工。", "私人、自然人（未统一）"),
        ("ENT-008", "其他", "otherCounterparty", "对手方类型", "无法归入客户、供应商、员工、个人或已知关联方的对手方。", "不得作为长期未治理分类，应支持后续补充主档。", "未知（带判断含义）"),
        ("ENT-009", "已知关联方", "knownRelatedParty", "主体", "关联方主档已维护且在当前分析期间有效的关联主体。", "以关联方主档和关系生效期间为准。", "高风险关联方"),
        ("ENT-010", "潜在关联主体", "potentialRelatedParty", "分析对象", "基于工商、账户、地址、人员或交易行为线索识别，尚待项目组确认的主体。", "只能表述为线索，不直接等同于已知关联方。", "关联方（未经确认）"),
        ("ENT-011", "被审计单位主档", "auditEntityMaster", "主档", "记录被审计单位及集团范围公司身份、关系和有效期间的标准清单。", "关键属性：公司名称、统一社会信用代码、关系、持股比例、有效期间。", "公司清单"),
        ("ENT-012", "关联方主档", "relatedPartyMaster", "主档", "记录已识别关联方、关联关系及生效终止期间的标准清单。", "分析时按期间有效性匹配；外部识别结果需经补充后进入主档。", "关联方清单（未说明主档）"),
        ("ENT-013", "银行账号主档", "bankAccountMaster", "主档", "记录被审计单位银行账号、开户行、币种、启用和停用期间的标准清单。", "用于账户完整性、账号范围和同名账户校验。", "银行账户清单"),
        ("ENT-014", "客户主档", "customerMaster", "主档", "记录客户身份、类别、批准及无效期间的标准清单。", "用于客户识别、主档匹配和期间有效性判断。", "客户清单"),
        ("ENT-015", "供应商主档", "supplierMaster", "主档", "记录供应商身份、类别、批准及无效期间的标准清单。", "用于供应商识别、主档匹配和期间有效性判断。", "供应商清单"),
        ("ENT-016", "员工主档", "employeeMaster", "主档", "记录员工身份、职位、入离职期间及账户信息的标准清单。", "用于员工交易、个人交易和有效任职期间判断。", "员工清单"),
        ("ENT-017", "统一社会信用代码", "unifiedSocialCreditCode", "身份属性", "境内法人和其他组织的统一身份识别代码。", "优先作为公司实体去重键；缺失时才使用标准化名称辅助匹配。", "工商注册号（混用）"),
        ("ENT-018", "税号", "taxpayerId", "身份属性", "发票购销双方的纳税人识别号。", "企业通常与统一社会信用代码一致，但系统仍保留原始税号字段。", "唯一识别码（含义不清）"),
        ("ENT-019", "外部工商数据", "externalBusinessData", "数据源", "由天眼查等外部工商数据源获得的公司和股权关系清单。", "仅用于范围核对和关系线索；最终分析范围以确认后的主档为准。", "天眼查主档"),
        ("ENT-020", "ARMS来源", "armsSource", "数据源", "来自ARMS系统的被审计单位或关系主体清单来源标识。", "在分析范围核对中与工商数据和主档并列展示。", "ARMS主档"),
    ]),
    ("03", "资金流水与账户", "统一银行流水数据对象、金额方向和账户校验口径。", [
        ("BNK-001", "银行流水", "bankTransaction", "数据对象", "银行提供的账户收付款及余额变动原始记录。", "关键属性：交易日期、交易时间、本方账号、对方名称、方向、金额、余额、流水识别号。", "资金流水（指原始数据时）"),
        ("BNK-002", "流水明细", "transactionDetail", "记录", "标准化后的单笔银行流水记录。", "正式表格页签和导出字段名称；应保留唯一流水识别号。", "交易明细、流水详情"),
        ("BNK-003", "流水识别号", "transactionId", "唯一标识", "用于唯一识别银行流水记录的系统或银行来源编号。", "去重、穿透和规则命中关联必须使用；为空时生成稳定替代键。", "流水号（可能与银行凭证号混淆）"),
        ("BNK-004", "本方名称", "ownPartyName", "字段", "银行流水中账户归属的被审计单位名称。", "需与被审计单位主档标准化匹配。", "本公司名称"),
        ("BNK-005", "本方账号", "ownAccount", "字段", "被审计单位发生交易的银行账号。", "全平台统一使用“账号”，不使用“帐号”。", "本方帐号、银行账号（未说明方向）"),
        ("BNK-006", "对方名称", "counterpartyName", "字段", "银行流水中记录的交易对手名称。", "显示原始值时标注“原始名称”；分析字段使用标准化对手方名称。", "对手名称"),
        ("BNK-007", "流入金额", "inflowAmount", "金额指标", "资金进入本方银行账户的交易金额。", "按当前单位切换器缩放展示；金额为非负数。", "累计流入（明细列）、收入金额"),
        ("BNK-008", "流出金额", "outflowAmount", "金额指标", "资金离开本方银行账户的交易金额。", "按当前单位切换器缩放展示；金额为非负数。", "累计流出（明细列）、支出金额"),
        ("BNK-009", "交易总额", "grossTransactionAmount", "金额指标", "流入金额与流出金额之和。", "公式：流入金额＋流出金额；不得与净流量混用。", "流水总额（未定义）"),
        ("BNK-010", "净资金流入", "netCashInflow", "金额指标", "流入金额扣除流出金额后的净额。", "公式：流入金额－流出金额；可为负数。", "净流量（未说明方向）"),
        ("BNK-011", "交易后余额", "postTransactionBalance", "金额字段", "单笔交易入账后银行账户记录的余额。", "应与交易日期、交易时间和流水顺序共同使用。", "余额（未说明时点）"),
        ("BNK-012", "等值人民币余额", "cnyEquivalentBalance", "金额字段", "外币余额按指定汇率折算后的人民币金额。", "需保留原币币种、原币金额、折算汇率和汇率日期。", "人民币余额（未说明折算）"),
        ("BNK-013", "资金变动趋势", "fundMovementTrend", "分析模块", "展示银行账户余额及流入流出随时间的变化。", "横轴为日期，纵轴金额随顶部单位切换；支持年度范围选择。", "资金余额变动趋势"),
        ("BNK-014", "资金总览", "fundOverview", "分析模块", "按被审计单位或账户汇总资金规模、笔数、对手方和文件覆盖。", "合计状态可展开年度明细；不使用固定单位后缀。", "账户资金总览"),
        ("BNK-015", "余额连续性校验", "balanceContinuityCheck", "数据校验", "检查相邻流水的交易后余额与前后交易金额是否能够连续衔接。", "结果区分当日余额不连续、无流水数据、校验无误和余额错误。", "余额连续性判断"),
        ("BNK-016", "缺失同名账户", "missingSameNameAccount", "数据校验", "被审计单位同名银行账户出现在流水或外部材料中，但未在账号主档完整维护的情况。", "应记录缺失账号、金额、交易次数、缺失原因和支持性文件索引。", "缺失账户"),
        ("BNK-017", "总账核对", "generalLedgerReconciliation", "数据校验", "按公司和期间核对银行流水汇总与科目余额表总账数据。", "比较期初、借方、贷方、期末及其差异；保留a1至d2口径标识。", "总帐核对"),
        ("BNK-018", "分账号核对", "accountReconciliation", "数据校验", "按单个银行账号核对流水汇总与分账号科目余额表数据。", "除总账核对字段外，还包括银行账号、开户银行和数据来源。", "分账户核对"),
        ("BNK-019", "含集团内往来", "includingIntraGroup", "分析口径", "资金金额包含被审计单位集团内公司之间的往来。", "与“剔除集团内往来”并列展示，来源于被审计单位主档。", "集团口径（含义不清）"),
        ("BNK-020", "剔除集团内往来", "excludingIntraGroup", "分析口径", "从资金金额中剔除集团内公司之间的往来。", "仅剔除已确认集团范围内主体，不剔除潜在关联主体。", "不含关联方（范围错误）"),
    ]),
    ("04", "增值税发票与税务", "统一进销项方向、金额、税额和税表核对表达。", [
        ("VAT-001", "增值税发票", "vatInvoice", "数据对象", "记录购销双方、开票日期、金额、税额和发票状态的税务凭证。", "分销项发票和进项发票；正式标题不使用“税票”。", "税票"),
        ("VAT-002", "销项发票", "salesInvoice", "数据对象", "被审计单位作为销售方开具的增值税发票。", "主要对应银行流入；不代表已收款。", "销售发票（口径未统一）"),
        ("VAT-003", "进项发票", "purchaseInvoice", "数据对象", "被审计单位作为购买方取得的增值税发票。", "主要对应银行流出；不代表已付款。", "采购发票（口径未统一）"),
        ("VAT-004", "发票号码", "invoiceNumber", "唯一标识", "税务系统赋予单张发票的号码或号码组合。", "与发票代码、数电票号码等共同形成稳定唯一键。", "票号（正式字段）"),
        ("VAT-005", "销项金额", "salesInvoiceNetAmount", "金额指标", "销项发票不含税金额合计。", "不包含销项税额；按发票状态和当前筛选范围汇总。", "销项税金额"),
        ("VAT-006", "进项金额", "purchaseInvoiceNetAmount", "金额指标", "进项发票不含税金额合计。", "不包含进项税额；按发票状态和当前筛选范围汇总。", "进项税金额"),
        ("VAT-007", "销项税额", "salesVatAmount", "税额指标", "销项发票税额合计。", "不与销项金额或销项税价合计混用。", "销项税（表示金额时）"),
        ("VAT-008", "进项税额", "purchaseVatAmount", "税额指标", "进项发票税额合计。", "不与进项金额或进项税价合计混用。", "进项税（表示金额时）"),
        ("VAT-009", "销项税价合计", "salesVatTotal", "金额指标", "销项发票不含税金额与税额之和。", "公式：销项金额＋销项税额；全平台固定使用“税价”。", "销项价税合计、销项税价总额"),
        ("VAT-010", "进项税价合计", "purchaseVatTotal", "金额指标", "进项发票不含税金额与税额之和。", "公式：进项金额＋进项税额；全平台固定使用“税价”。", "进项价税合计、进项税价总额"),
        ("VAT-011", "销项税率", "salesVatRate", "比例指标", "销项发票适用税率。", "多税率汇总时不得简单平均，应按不含税金额加权或分税率展示。", "销项发票税率（可作为说明）"),
        ("VAT-012", "进项税率", "purchaseVatRate", "比例指标", "进项发票适用税率。", "多税率汇总时不得简单平均，应按不含税金额加权或分税率展示。", "进项发票税率（可作为说明）"),
        ("VAT-013", "购销趋势", "purchaseSalesTrend", "分析模块", "按月份展示销项金额和进项金额的变化。", "默认使用不含税金额；图例应明确销项、进项。", "进销项趋势"),
        ("VAT-014", "购销税额构成", "purchaseSalesTaxComposition", "分析模块", "按对手方类型或关系分类展示销项、进项金额及税额构成。", "支持销项税、进项税、销项税VS进项税三种视角。", "购销构成"),
        ("VAT-015", "缺失同名购销方", "missingSameNameCounterparty", "数据校验", "发票购销方名称与被审计单位或主档名称匹配，但对应税号未完整维护的情况。", "应记录缺失税号、金额、税额、税价合计、原因和索引。", "缺失购销方"),
        ("VAT-016", "发票状态", "invoiceStatus", "状态属性", "发票当前有效状态。", "枚举至少包括正常、作废、红冲；分析口径须明确是否纳入。", "票据状态"),
        ("VAT-017", "税表核对", "taxReturnReconciliation", "数据校验", "将增值税发票汇总与增值税申报表对应栏次进行核对。", "需按期间、方向和税率口径比较，并记录差异说明。", "申报表核对"),
    ]),
    ("05", "账期、匹配与综合分析", "定义资金与发票之间的时间归属、匹配和差异口径。", [
        ("CMP-001", "银行流水与增值税对比分析", "bankVatComparison", "分析页", "按总量、月度、账期、对手方构成和明细比较银行流水与增值税发票。", "客观分析，不自动输出异常或风险判断。", "银行流水与税票对比"),
        ("CMP-002", "账期规则设置", "periodRuleSettings", "规则入口", "设置全局默认账期口径并管理对手方单独账期规则。", "单独规则优先，未设置对手方使用默认口径。", "账期设置"),
        ("CMP-003", "默认账期口径", "defaultPeriodBasis", "规则", "未配置单独账期规则的对手方统一使用的税票月份归属规则。", "枚举：同期、税票前移N月、税票后移N月。", "默认账期规则"),
        ("CMP-004", "单独账期规则", "counterpartyPeriodRule", "规则", "为指定对手方单独配置的税票月份归属规则。", "优先级高于默认账期口径；应记录设置人和更新时间。", "特殊账期"),
        ("CMP-005", "同期", "samePeriod", "账期口径", "税票按原始开票月份参与银行流水月度对比。", "银行流水月份始终不平移。", "不平移"),
        ("CMP-006", "税票前移N月", "vatShiftEarlier", "账期口径", "将税票归属月份向前移动N个月后参与对比。", "例如原6月税票前移1月后归入5月。", "付款前移、账期修正"),
        ("CMP-007", "税票后移N月", "vatShiftLater", "账期口径", "将税票归属月份向后移动N个月后参与对比。", "例如原4月税票后移1月后归入5月。", "付款后移、账期修复"),
        ("CMP-008", "流入票差", "inflowInvoiceDifference", "差异指标", "银行流入金额与销项税价合计之间的差额。", "公式：流入金额－销项税价合计；保持中性展示。", "流入链路对比金额"),
        ("CMP-009", "流出票差", "outflowInvoiceDifference", "差异指标", "银行流出金额与进项税价合计之间的差额。", "公式：流出金额－进项税价合计；保持中性展示。", "流出链路对比金额"),
        ("CMP-010", "资金—票流匹配", "fundInvoiceMatching", "分析关系", "将银行收付款与发票按对手方、方向、金额和时间窗口建立对应关系。", "支持一对一、一对多和多对一；每组对应关系应有唯一编号。", "票据资金对应（未定义）"),
        ("CMP-011", "匹配状态", "matchingStatus", "状态", "资金与发票当前对应结果。", "枚举：已匹配、部分匹配、有资金无发票、有发票无资金、待确认。", "对应状态"),
        ("CMP-012", "匹配金额", "matchedAmount", "金额指标", "已建立资金—票流关系的金额。", "一对多、多对一场景不得超过任一侧可分配余额。", "对应金额"),
        ("CMP-013", "日期差异", "dateDifference", "期间指标", "银行交易日期与发票日期之间的天数或月份差。", "需明确正负方向；页面账期平移按月份处理。", "时间差（未说明方向）"),
        ("CMP-014", "同月跨年对比", "sameMonthYearComparison", "图表模式", "将多个年度相同月份的数据进行并列比较。", "支持选择1至3年；业务指标用色相，年度用深浅。", "同年跨月对比"),
        ("CMP-015", "时序趋势", "timeSeriesTrend", "图表模式", "按连续时间顺序展示指标变化。", "可使用柱状图或折线图；月份与年度按实际时间排列。", "同年跨月趋势"),
        ("CMP-016", "分析时间", "analysisTime", "筛选条件", "当前分析采用的起止月份或日期范围。", "筛选器统一使用；报告中可称“分析期间”。", "月份范围（作为通用字段）"),
        ("CMP-017", "三年一期", "threeYearsCurrentPeriod", "期间口径", "三个完整年度与当前一期数据的组合分析口径。", "页面表头直接显示年度和当前期间，不在指标名称中反复写“三年一期”。", "多年多期（未具体）"),
    ]),
    ("06", "数据校验、特殊交易与报告", "统一质量问题、规则命中、风险线索和成果导出术语。", [
        ("CHK-001", "数据问题", "dataIssue", "质量状态", "影响分析完整性、准确性或可比性的上传数据或主档问题。", "例如未认领、账号主档不完整、余额不连续；应提供跳转处理入口。", "风险（数据质量问题场景）"),
        ("CHK-002", "分析范围清单核对", "analysisScopeReconciliation", "范围校验", "对比外部工商数据、ARMS来源与被审计单位主档，识别公司范围遗漏。", "三方共同包含才视为核对一致；缺失时可补充至主档。", "公司主档遗漏检查"),
        ("CHK-003", "账号范围分析", "accountScopeAnalysis", "范围校验", "对比银行账号主档、科目余额表和缺失账号清单的账户覆盖。", "仅展示范围与备注，不直接执行纳入或剔除。", "账号清单检查"),
        ("CHK-004", "特殊交易识别", "specialTransactionDetection", "分析页", "从全部交易中按预设核查标准识别需要进一步审计关注的交易线索。", "不以对手方是否已标记为高风险为筛查前提。", "高风险核查"),
        ("CHK-005", "核查程序", "auditProcedure", "规则集合", "围绕特定审计目的组织的一组核查标准、参数和结果。", "主要分定量、定性、重点关注异常、工商信息异常和其他异常。", "核查项目"),
        ("CHK-006", "核查标准", "auditCriterion", "规则", "用于判断单笔交易或对手方是否命中某一核查条件的具体规则。", "应有唯一编号、名称、参数、适用对象和版本。", "异常规则（过度判断）"),
        ("CHK-007", "命中个数", "hitCount", "数量指标", "当前对象命中的核查标准数量。", "按对象去重后统计标准数，不等于交易笔数。", "风险个数"),
        ("CHK-008", "规则命中次数", "ruleHitCount", "数量指标", "所有交易对各规则的命中记录总数。", "同一交易命中多项规则时重复计数。", "高风险交易笔数"),
        ("CHK-009", "去重交易笔数", "uniqueTransactionCount", "数量指标", "按唯一流水识别号去重后的交易数量。", "用于高风险交易总量；不得与规则命中次数混用。", "命中次数"),
        ("CHK-010", "异常对手方", "exceptionCounterparty", "分析对象", "至少命中一项已定义核查标准或存在需解释差异的对手方。", "名称保持线索属性；不得直接表述为“高风险对手方”。", "高风险对手方"),
        ("CHK-011", "收支重合", "inflowOutflowOverlap", "交易特征", "同一对手方在分析期间同时存在资金流入和流出。", "应分别展示流入、流出金额及占比，并结合业务实质分析。", "同进同出、客商重合（泛化）"),
        ("CHK-012", "集中度", "concentration", "比例指标", "前N大对象金额占对应总金额的比例。", "应明确方向、N值、分母范围和是否剔除集团内往来。", "TOP占比（未说明口径）"),
        ("CHK-013", "黑名单", "fraudCollusionList", "外部清单", "AQPP-Audit提供的可能配合实施财务舞弊的第三方清单。", "命中仅作为审计线索；表格中使用勾选标识并支持穿透。", "造假主体（结论性）"),
        ("CHK-014", "重要性水平", "materiality", "审计参数", "项目组设定、用于抽样和审计判断的金额阈值。", "由工作簿维护并带入报告及底稿导出；单位为元。", "Materiality（仅英文）"),
        ("CHK-015", "信息测试 - 完整性测试", "completenessTest", "底稿程序", "验证上传数据是否完整覆盖应取得的账户、期间和记录。", "输出完整性测试结果及支持性说明。", "完整性测试（脱离上下文）"),
        ("CHK-016", "信息测试 - 准确性测试", "accuracyTest", "底稿程序", "通过直接测试或属性抽样评价流水记录的准确性。", "选择“否”时隐藏风险等级、抽样说明、样本量表格和分页。", "准确性测试（脱离上下文）"),
        ("CHK-017", "风险评估报告", "riskAssessmentReport", "输出成果", "汇总数据覆盖、资金发票画像、特殊交易和审计关注事项的正式报告。", "报告指标应与平台同一数据快照、单位和规则版本一致。", "分析报告（泛称）"),
        ("CHK-018", "资金流水核查底稿", "fundFlowWorkingPaper", "输出成果", "记录核查程序、参数、结果、抽样和项目组解释的审计底稿。", "应保留工作簿、数据、规则和导出时间信息。", "银行流水底稿"),
    ]),
    ("07", "状态与操作", "统一页面状态、数据状态和用户操作的动词表达。", [
        ("STA-001", "已包含", "included", "来源状态", "当前对象已存在于指定来源清单。", "使用绿色空心圆勾；不使用大面积绿色底色。", "有、存在"),
        ("STA-002", "未包含", "notIncluded", "来源状态", "当前对象未存在于指定来源清单。", "使用红色空心圆减号；不直接表示异常结论。", "没有、缺失（未说明来源）"),
        ("STA-003", "核对一致", "reconciled", "核对状态", "参与比较的所有必要来源均包含该对象，或数值差异为零。", "需根据具体核对场景定义必要来源。", "正常"),
        ("STA-004", "待补充至主档", "pendingMasterSupplement", "处理状态", "外部来源存在但目标主档尚未维护。", "提供“补充至主档”操作；操作后记录不消失。", "缺失（无后续动作）"),
        ("STA-005", "已补充至主档", "supplementedToMaster", "处理状态", "用户已执行补充操作且当前行主档状态已更新。", "保留原行和操作痕迹，并更新统计数量。", "已处理（含义过宽）"),
        ("STA-006", "待核实", "pendingVerification", "处理状态", "来源信息存在冲突或证据不足，需要项目组进一步确认。", "不得自动纳入或剔除分析范围。", "异常"),
        ("STA-007", "本期新增", "newInCurrentPeriod", "跨期状态", "当前期间首次出现，历史比较期间未出现。", "需基于同一实体键和完整历史期间判断。", "新增（未说明期间）"),
        ("STA-008", "持续存在", "persistent", "跨期状态", "相同对象或事项连续两个及以上期间出现。", "连续期间数应可追溯。", "一直异常"),
        ("STA-009", "历史复发", "recurring", "跨期状态", "历史期间出现、随后消失并在当前期间再次出现。", "与持续存在区分。", "重复发生（未区分连续性）"),
        ("STA-010", "风险缓释", "riskMitigated", "跨期状态", "风险指标或未解释事项较历史期间下降并有支持性证据。", "仅表示变化，不代表风险已消除。", "风险消失"),
        ("ACT-001", "查询", "query", "操作", "按当前筛选条件刷新结果。", "按钮文字固定；不得与“搜索”混用。", "搜索、确认查询"),
        ("ACT-002", "重置", "reset", "操作", "将当前页面筛选恢复为默认值。", "不删除已保存备注、主档规则或处理结果。", "清空（除非仅清空输入）"),
        ("ACT-003", "自定义表头", "customizeColumns", "操作", "控制表格字段显示与隐藏。", "勾选后即时生效，不需要确认按钮。", "列设置、自定义列"),
        ("ACT-004", "导出清单", "exportList", "操作", "导出当前表格在当前筛选口径下的数据。", "文件应包含筛选摘要、单位、数据版本和导出时间。", "导出数据"),
        ("ACT-005", "导出结果", "exportResult", "操作", "导出当前核对或分析结果。", "用于范围核对等结果型卡片；与“导出清单”按对象区分。", "下载结果"),
        ("ACT-006", "查看明细", "viewDetails", "操作", "打开当前对象的明细抽屉或弹窗。", "表格操作统一使用；右侧抽屉优先。", "查看详情、详情"),
        ("ACT-007", "展开 / 收起", "expandCollapse", "操作", "切换卡片、年度列组或分类列组的可见状态。", "卡片使用圆形箭头按钮；不改变数据筛选状态。", "打开 / 关闭"),
        ("ACT-008", "补充至主档", "supplementToMaster", "操作", "将当前外部识别对象补充到目标主档。", "操作完成后显示“已补充至主档”，原记录保留。", "加入清单"),
        ("ACT-009", "生成分析结果", "generateAnalysis", "操作", "按当前数据和参数重新生成分析结果。", "属于计算动作，不用于普通查询。", "生成结果"),
        ("ACT-010", "生成流水", "generateTransactions", "操作", "将上传文件标准化生成可分析的银行流水数据。", "生成后应记录数据版本和处理时间。", "生成明细"),
    ]),
]


ATTRIBUTE_ROWS = [
    ("项目", "项目编号", "projectId", "字符串", "项目的唯一业务编号。", "必填；全平台不可重复。"),
    ("项目", "项目名称", "projectName", "字符串", "审计项目的正式名称。", "必填；按完整名称展示。"),
    ("工作簿", "工作簿名称", "workbookName", "字符串", "项目下分析任务的名称。", "必填；项目内唯一。"),
    ("工作簿", "分析期间", "analysisPeriod", "月份区间", "数据分析采用的起止月份。", "必填；格式YYYY-MM至YYYY-MM。"),
    ("工作簿", "报告期间", "reportingPeriod", "日期区间", "财务报告或审计报告覆盖期间。", "必填；与分析期间可不同。"),
    ("工作簿", "重要性水平", "materiality", "数值 / 元", "用于抽样和审计判断的金额阈值。", "可空；为空显示“—”，不得默认为0。"),
    ("被审计单位", "公司名称", "auditEntityName", "字符串", "被审计单位主档中的标准公司名称。", "必填；不省略展示。"),
    ("被审计单位", "统一社会信用代码", "unifiedSocialCreditCode", "字符串 / 18位", "公司实体的主要去重和匹配键。", "缺失时标记待补充，不用空字符串替代。"),
    ("被审计单位", "集团关系", "groupRelationship", "枚举", "该公司与集团或发行人的关系。", "如母公司、控股子公司、参股公司、分支机构。"),
    ("被审计单位", "持股比例", "holdingPercentage", "百分比", "直接或最终持股比例。", "未知显示“—”；0与未知必须区分。"),
    ("银行账户", "本方账号", "ownAccount", "字符串", "被审计单位银行账号。", "必填；界面可脱敏，导出按权限控制。"),
    ("银行账户", "开户银行", "bankName", "字符串", "银行账户开户机构名称。", "缺失显示“—”，进入主档完整性问题。"),
    ("银行账户", "币种", "currency", "ISO币种代码", "账户和交易原币币种。", "默认不推定人民币；使用CNY/RMB需全平台统一。"),
    ("银行账户", "启用日期", "accountOpenDate", "日期", "账户开始有效或开户日期。", "未知为空，不使用分析期起始日替代。"),
    ("银行账户", "停用日期", "accountCloseDate", "日期", "账户注销或停止使用日期。", "有效账户为空；注销后交易需单独识别。"),
    ("银行流水", "交易日期", "transactionDate", "日期", "银行交易入账日期。", "必填；格式YYYY-MM-DD。"),
    ("银行流水", "交易时间", "transactionTime", "时间", "银行记录的交易时刻。", "可空；为空不参与异常时点规则。"),
    ("银行流水", "对方名称", "counterpartyRawName", "字符串", "银行原始数据中的对方户名。", "保留原值；另存标准化名称。"),
    ("银行流水", "收付方向", "direction", "枚举", "交易相对本方账户的资金方向。", "仅允许流入、流出；未知进入数据问题。"),
    ("银行流水", "流入金额", "inflowAmount", "数值 / 动态单位", "方向为流入的交易金额。", "无流入显示0；单位由顶部切换器控制。"),
    ("银行流水", "流出金额", "outflowAmount", "数值 / 动态单位", "方向为流出的交易金额。", "无流出显示0；单位由顶部切换器控制。"),
    ("银行流水", "交易后余额", "postTransactionBalance", "数值 / 动态单位", "该笔交易后的账户余额。", "源数据缺失则为空，不显示0。"),
    ("银行流水", "流水识别号", "transactionId", "字符串", "单笔交易的稳定唯一键。", "必填；源数据缺失时生成稳定替代键并标记来源。"),
    ("银行流水", "摘要", "transactionSummary", "字符串", "银行提供的交易摘要或用途。", "保留原文；空值显示“—”。"),
    ("增值税发票", "发票号码", "invoiceNumber", "字符串", "单张发票的识别号码。", "与发票代码等组成唯一键；不得转数值。"),
    ("增值税发票", "发票方向", "invoiceDirection", "枚举", "发票相对被审计单位的进销项方向。", "仅允许销项、进项。"),
    ("增值税发票", "开票日期", "invoiceDate", "日期", "发票开具日期。", "必填；账期平移基于其原始月份。"),
    ("增值税发票", "不含税金额", "netAmount", "数值 / 动态单位", "发票不含增值税的金额。", "缺失为空；不以税价合计倒推，除非规则明确。"),
    ("增值税发票", "税额", "taxAmount", "数值 / 动态单位", "发票列示的增值税税额。", "缺失为空；零税率发票可为0。"),
    ("增值税发票", "税价合计", "grossAmount", "数值 / 动态单位", "不含税金额与税额之和。", "公式校验失败时保留原值并记录差异。"),
    ("增值税发票", "税率", "taxRate", "百分比", "发票适用增值税税率。", "免税或不征税应使用状态枚举，不混作0%。"),
    ("增值税发票", "发票状态", "invoiceStatus", "枚举", "发票当前有效状态。", "正常、作废、红冲；空值进入数据问题。"),
    ("对手方", "标准名称", "counterpartyName", "字符串", "实体解析后的对手方统一名称。", "为空时回退显示原始名称并标记待治理。"),
    ("对手方", "对手方类型", "counterpartyType", "枚举", "对手方业务身份分类。", "客户、供应商、员工、个人、其他；关联方为关系属性。"),
    ("对手方", "关联关系", "relatedPartyRelationship", "枚举", "对手方与被审计单位的关联身份。", "非关联方为空；不得填“普通”。"),
    ("对手方", "所属集团", "groupName", "字符串", "对手方所属企业集团标准名称。", "未知为空；不得使用“其他”代替。"),
    ("对手方", "黑名单命中", "blacklistHit", "布尔", "是否命中第三方配合造假清单。", "命中显示勾，未命中留空或显示减号。"),
    ("账期规则", "规则类型", "periodRuleType", "枚举", "税票月份归属方式。", "统一使用同期、税票前移、税票后移。"),
    ("账期规则", "平移月数", "shiftMonths", "整数 / 月", "税票月份平移的绝对月数。", "同期为0；前移后移至少1。"),
    ("账期规则", "规则来源", "periodRuleSource", "枚举", "当前生效规则的来源。", "默认口径或单独规则。"),
    ("资金—票流匹配", "匹配关系编号", "matchId", "字符串", "一组资金与发票对应关系的唯一编号。", "必须填写；支持一对一、一对多、多对一。"),
    ("资金—票流匹配", "匹配状态", "matchingStatus", "枚举", "当前资金与发票对应状态。", "已匹配、部分匹配、有资金无发票、有发票无资金、待确认。"),
    ("资金—票流匹配", "匹配金额", "matchedAmount", "数值 / 动态单位", "已分配到匹配关系的金额。", "未匹配显示0；不得超过可匹配余额。"),
    ("资金—票流匹配", "日期差异", "dateDifferenceDays", "整数 / 天", "银行交易日期减发票日期。", "正值表示资金晚于发票；空值表示任一侧日期缺失。"),
    ("核查结果", "核查标准编号", "criterionCode", "字符串", "核查标准的稳定编号。", "必填；规则升级不得复用旧编号含义。"),
    ("核查结果", "命中标识", "hitFlag", "布尔", "当前对象是否满足指定核查标准。", "命中显示勾，未命中显示空白或减号。"),
    ("核查结果", "命中个数", "hitCount", "整数 / 个", "当前对象实际命中的标准数量。", "无命中显示0，不显示“—”。"),
    ("核查结果", "首次发生期间", "firstOccurrencePeriod", "月份或年度", "该风险线索首次出现的完整期间。", "无历史数据时为空并注明历史覆盖不足。"),
    ("核查结果", "跨期状态", "crossPeriodStatus", "枚举", "事项在多年多期中的变化状态。", "本期新增、持续存在、风险上升、历史复发、风险缓释。"),
    ("通用", "备注", "remark", "字符串", "项目组对对象、差异或处理结果的说明。", "不参与计算；记录更新人和更新时间。"),
    ("通用", "数据版本", "dataVersion", "字符串", "当前分析所使用的标准化数据快照版本。", "导出成果必填。"),
    ("通用", "规则版本", "ruleVersion", "字符串", "当前核查规则集合版本。", "特殊交易结果和报告必填。"),
    ("通用", "更新时间", "updatedAt", "日期时间", "记录最近一次更新时点。", "格式YYYY-MM-DD HH:mm；使用项目时区。"),
    ("通用", "更新人", "updatedBy", "用户标识", "最近一次修改记录的用户。", "自动记录，不允许手工输入。"),
]


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.38)
    section.footer_distance = Inches(0.38)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_CN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15
    for style_name, size, before, after in (("Heading 1", 15, 14, 7), ("Heading 2", 12, 10, 5), ("Heading 3", 10.5, 8, 4)):
        style = styles[style_name]
        style.font.name = FONT_CN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(DARK_BLUE if style_name == "Heading 1" else BLUE)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    p = header.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    left = p.add_run("AUDIT COMPASS  |  系统统一名词与属性定义表")
    set_run_font(left, 8, bold=True, color=MID_BLUE)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:color"), BORDER)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)
    footer = section.footer
    fp = footer.paragraphs[0]
    add_page_number(fp)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(14)
    title.paragraph_format.space_after = Pt(4)
    tr = title.add_run("Audit Compass 系统统一名词与属性定义表")
    set_run_font(tr, 22, bold=True, color=DARK_BLUE)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    sr = subtitle.add_run("统一平台各类名称、定义和主要属性口径")
    set_run_font(sr, 11, color=MID_BLUE)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(12)
    for label, value in (("版本", "V1.0"), ("生效日期", "2026-08-18"), ("适用范围", "产品需求、UI设计、前后端开发、测试、导出文件与分析报告")):
        lr = meta.add_run(f"{label}：")
        set_run_font(lr, 9, bold=True, color=MUTED)
        vr = meta.add_run(value + "    ")
        set_run_font(vr, 9, color=INK)

    add_callout(doc, "文档定位", "本表用于统一平台页面、需求、测试、导出文件和分析报告中的名称及定义。出现不同表达时，以本表中的标准名称和口径说明为准。")

    add_section_heading(doc, "00", "使用说明", "按业务主题查找名称，并统一用于页面、需求、测试和报告。")
    intro_rows = [
        ("名称层级", "平台 → 项目 → 工作簿 → 功能模块 → Tab → Card → 字段或指标 → 状态或操作。"),
        ("使用顺序", "先确认名称表达的内容和计算口径，再用于页面、表格、导出文件和报告。"),
        ("单位规则", "金额单位由平台顶部单位切换器统一控制，字段名称和单元格数值不固定附加K、M或“千元”等单位。"),
        ("判断词规则", "“异常、风险、预警”等词仅用于已定义核查标准的结果；客观对比页面使用差异、匹配、覆盖、构成等中性名称。"),
        ("变更规则", "新增或修改名称时，先更新本表，再同步页面、需求、测试、导出文件和报告。"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [1850, 7510])
    set_table_borders(table)
    for i, header_text in enumerate(("规则", "统一要求")):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, LIGHT_BLUE)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header_text)
        set_run_font(run, 8.5, bold=True, color=DARK_BLUE)
    for idx, (label, text) in enumerate(intro_rows):
        row = table.add_row()
        prevent_row_split(row)
        if idx % 2:
            set_cell_shading(row.cells[0], "FBFCFE")
            set_cell_shading(row.cells[1], "FBFCFE")
        for cell in row.cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        r1 = row.cells[0].paragraphs[0].add_run(label)
        set_run_font(r1, 8.5, bold=True, color=INK)
        r2 = row.cells[1].paragraphs[0].add_run(text)
        set_run_font(r2, 8.5, color=INK)

    for number, title_text, description, rows in TERM_SECTIONS:
        add_section_heading(doc, number, title_text, description)
        add_term_table(doc, rows)

    add_section_heading(doc, "08", "主要属性名称与定义", "汇总平台常用字段的名称、定义和使用说明。")
    add_attribute_table(doc, ATTRIBUTE_ROWS)

    add_callout(doc, "使用原则", "页面、需求、测试、导出文件和报告应使用本表中的统一名称；如需新增名称，应同时补充定义和使用说明。")

    props = doc.core_properties
    props.title = "Audit Compass 系统统一名词与属性定义表"
    props.subject = "平台统一名称、业务定义与主要属性"
    props.author = "Audit Compass 产品与数据治理"
    props.keywords = "统一名称, 名词定义, 审计分析, 银行流水, 增值税发票"
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
